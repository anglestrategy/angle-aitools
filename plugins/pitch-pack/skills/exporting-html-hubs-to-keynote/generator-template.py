#!/usr/bin/env python3
"""
[CLIENT] [PROJECT] — Strategy Approach Keynote/PPTX generator (template)

Replace the token block below with the client's design system values, then
build the slides by calling the slide_* helpers in BUILD section at the bottom.

Output: "[Client] - Strategy Approach.pptx", 1920×1080, native editable.

Open in Keynote: File → Open → select the .pptx
Or export to .key: File → Export To → Keynote

────────────────────────────────────────────────────────────
TOKENS — replace with the client design system
────────────────────────────────────────────────────────────
"""

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree
from pathlib import Path

# ════════════════════════════════════════════════════════════
# CLIENT TOKENS — edit these
# ════════════════════════════════════════════════════════════

SLIDE_W = 1920
SLIDE_H = 1080

C_SURFACE_LIGHT = RGBColor(0xDB, 0xE2, 0xE0)   # OM Gray — replace
C_SURFACE_DARK  = RGBColor(0x00, 0x00, 0x00)
C_TEXT_LIGHT    = RGBColor(0x00, 0x00, 0x00)
C_TEXT_DARK     = RGBColor(0xFF, 0xFF, 0xFF)
C_ACCENT        = RGBColor(0x00, 0xFF, 0x00)   # OM Green — replace
C_ACCENT_HEX    = '00FF00'

FONT = 'OM'                                     # client font name as installed

# cqi = % of slide width
def cqi(n):    return n / 100 * SLIDE_W
def cqi_pt(n): return cqi(n) * 0.75              # 96 dpi conversion

PAD_X = cqi(6.5)                                 # horizontal padding
PAD_Y = cqi(4.56)                                # vertical padding
INNER_W = SLIDE_W - 2 * PAD_X

# Font sizes (pt)
FS_COVER       = 100
FS_NUMBER      = 180
FS_TITLE       = round(cqi_pt(4.4))              # ≈63
FS_STATEMENT   = 56
FS_STATEMENT_T = 70
FS_HEAD_4      = 32
FS_LEAD        = 22
FS_BODY        = round(cqi_pt(1.3))              # ≈19
FS_DECK        = 18
FS_KICKER      = 11

# Hairline y-positions
HL_COVER_TOP = SLIDE_H * 0.154
HL_COVER_MID = SLIDE_H * 0.332
HL_COVER_BOT = SLIDE_H * 0.75
HL_DIV_TOP   = SLIDE_H * 0.154
HL_DIV_BOT   = SLIDE_H * 0.665

HAIRLINE_W = 2
SPC_KICKER = 200

# ════════════════════════════════════════════════════════════
# HELPERS
# ════════════════════════════════════════════════════════════

def px_emu(n): return int(round(n * 9525))

def add_slide(prs, fill):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = fill
    return s

def add_hairline(slide, y, color, full=True, weight=HAIRLINE_W):
    x1 = 0 if full else PAD_X
    x2 = SLIDE_W if full else SLIDE_W - PAD_X
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
        px_emu(x1), px_emu(y), px_emu(x2), px_emu(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln

def _new_textbox(slide, x, y, w, h, anchor='top'):
    tb = slide.shapes.add_textbox(px_emu(x), px_emu(y), px_emu(w), px_emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ('left','right','top','bottom'):
        setattr(tf, f'margin_{m}', 0)
    tf.vertical_anchor = {
        'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM
    }[anchor]
    return tf

def _style_run(run, size, color, bold=False, letter_spacing=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    if letter_spacing is not None:
        rPr = run._r.get_or_add_rPr()
        rPr.set('spc', str(letter_spacing))

def set_highlight(run, hex_color=C_ACCENT_HEX):
    """Inline OOXML highlight (renders in Keynote and recent PowerPoint)."""
    rPr = run._r.get_or_add_rPr()
    h = etree.SubElement(rPr, qn('a:highlight'))
    srgb = etree.SubElement(h, qn('a:srgbClr'))
    srgb.set('val', hex_color)

def text(slide, content, x, y, w, h, size, color,
         bold=False, anchor='top', align='left',
         line_spacing=None, letter_spacing=None):
    tf = _new_textbox(slide, x, y, w, h, anchor)
    p = tf.paragraphs[0]
    p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
    if line_spacing is not None: p.line_spacing = line_spacing
    run = p.add_run(); run.text = content
    _style_run(run, size, color, bold, letter_spacing)
    return tf

def rt(text_str, highlight=False, bold=False, color=None):
    return {'text': text_str, 'highlight': highlight, 'bold': bold, 'color': color}

def rich_text(slide, segments, x, y, w, h, size, color,
              bold=False, anchor='top', align='left',
              line_spacing=None, letter_spacing=None):
    tf = _new_textbox(slide, x, y, w, h, anchor)
    p = tf.paragraphs[0]
    p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
    if line_spacing is not None: p.line_spacing = line_spacing
    for seg in segments:
        run = p.add_run(); run.text = seg['text']
        if seg.get('highlight'):
            _style_run(run, size, C_TEXT_LIGHT, bold=True, letter_spacing=letter_spacing)
            set_highlight(run)
        else:
            _style_run(run, size, seg.get('color') or color,
                       bold=seg.get('bold') or bold, letter_spacing=letter_spacing)
    return tf

# ════════════════════════════════════════════════════════════
# SLIDE TEMPLATES
# (Add more as you need them: ring, arc, bench, objectives, filters.
#  See the worked example in build_strategy_keynote.py — Waha project.)
# ════════════════════════════════════════════════════════════

def slide_cover(prs, title_segments, top_left, top_right, strapline):
    s = add_slide(prs, C_SURFACE_LIGHT)
    add_hairline(s, HL_COVER_TOP, C_TEXT_LIGHT)
    add_hairline(s, HL_COVER_BOT, C_TEXT_LIGHT)
    text(s, top_left,  PAD_X, HL_COVER_TOP - 50, INNER_W * 0.6, 30,
         FS_KICKER, C_TEXT_LIGHT, letter_spacing=SPC_KICKER)
    text(s, top_right, PAD_X + INNER_W * 0.6, HL_COVER_TOP - 50, INNER_W * 0.4, 30,
         FS_KICKER, C_TEXT_LIGHT, align='right', letter_spacing=SPC_KICKER)
    rich_text(s, title_segments,
              PAD_X, HL_COVER_TOP + 100, INNER_W, 520,
              FS_COVER, C_TEXT_LIGHT, bold=True, line_spacing=0.96)
    text(s, strapline, PAD_X, HL_COVER_BOT + 35, INNER_W * 0.7, 200,
         FS_LEAD, C_TEXT_LIGHT, line_spacing=1.45)
    return s


def slide_divider(prs, num_str, kicker_left, kicker_right, title_segments, deck,
                  dark=False):
    fill = C_SURFACE_DARK if dark else C_SURFACE_LIGHT
    color = C_TEXT_DARK if dark else C_TEXT_LIGHT
    s = add_slide(prs, fill)
    add_hairline(s, HL_DIV_TOP, color)
    add_hairline(s, HL_DIV_BOT, color)
    text(s, kicker_left,  PAD_X, HL_DIV_TOP - 50, INNER_W * 0.5, 30,
         FS_KICKER, color, letter_spacing=SPC_KICKER)
    text(s, kicker_right, PAD_X + INNER_W * 0.5, HL_DIV_TOP - 50, INNER_W * 0.5, 30,
         FS_KICKER, color, align='right', letter_spacing=SPC_KICKER)
    text(s, num_str, PAD_X, HL_DIV_TOP + 40, INNER_W * 0.22, 480,
         FS_NUMBER, C_ACCENT, anchor='top', line_spacing=0.88)
    rich_text(s, title_segments,
              PAD_X + INNER_W * 0.24, HL_DIV_TOP + 90, INNER_W * 0.76, 480,
              FS_TITLE, color, bold=True, line_spacing=1.02)
    text(s, deck,
         PAD_X + INNER_W * 0.24, HL_DIV_BOT + 35, INNER_W * 0.65, 280,
         FS_DECK, color, line_spacing=1.5)
    return s

# Add slide_statement, slide_two_cards, slide_single_panel, slide_ring, slide_arc,
# slide_bench, slide_objectives, slide_filters as needed.
# Lift them from build_strategy_keynote.py (Waha worked example).

# ════════════════════════════════════════════════════════════
# BUILD — write one call per slide
# ════════════════════════════════════════════════════════════

def build():
    prs = Presentation()
    prs.slide_width  = Emu(px_emu(SLIDE_W))
    prs.slide_height = Emu(px_emu(SLIDE_H))

    slide_cover(prs,
        title_segments=[rt('Strategy '), rt('Approach.', highlight=True)],
        top_left='[AGENCY] × [CLIENT]',
        top_right='STRATEGY APPROACH / [PROJECT]',
        strapline='[Strapline copy.]')

    slide_divider(prs, '01',
        '01 · WHERE WE STAND', 'THE MOMENT',
        [rt('[Section title with '), rt('accent', highlight=True), rt(' word.]')],
        '[Deck paragraph.]',
        dark=False)

    # ... add the rest of the slides ...

    out = Path(__file__).parent / '[CLIENT] - Strategy Approach.pptx'
    prs.save(out)
    print(f'Saved: {out}')
    print(f'Slides: {len(prs.slides)}')
    return out


if __name__ == '__main__':
    build()
